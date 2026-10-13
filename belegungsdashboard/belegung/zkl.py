"""Zimmerkontrollliste (ZKL, Excel des Internatsdienstes): je Zimmer Kundengruppe, TN-ID (Rios), Name, von/bis,
Zimmer-Infos (Boden, Renovierung, Hinweise) und interne Anmerkungen.

Der Gebäudeplan bleibt die Grundlage – er ist meist neuer. Die ZKL ergänzt, was er nicht weiß:

* **Tier:** Kundengruppe „Reha-Tier“/„Reha mit Tier“ → die Person hat ein Tier (Regel Haus 3.1 UG).
* **TN-ID:** für Gebäudeplan-Bewohner (Mails mit TN-Nr. finden die Person).
* **Zimmer-Infos:** Boden, Renovierung, Hinweise und Anmerkungen – in der Zimmerkarte.

Der **Abgleich** zeigt Abweichungen zum Entscheiden:

* Laut ZKL bewohnt, im Gebäudeplan frei → als belegt übernehmen (oder: ist ausgezogen).
* Laut ZKL Lager/Büro/Appartment, im Zimmerplan belegbar → „gehört nicht zum Internat“.
* Laut ZKL Gästehaus (SGH) bzw. Tier-Pool, im Zimmerplan nicht → Merkmal setzen.
* Person in einem anderen Zimmer als laut ZKL (nur Info, der Gebäudeplan gilt).

Haus 3.2 heißt in der ZKL durchgehend „3.3“ – Zimmernummern werden gegen die Stammdaten abgeglichen.
Bei mehreren Zimmerblättern gilt das mit den neuesten Daten. Rein, ohne Qt. Gespeichert in ``importe/zkl.json``.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from . import speicher
from . import zimmerplan as zp
from .gebaeudeplan import Belegung

DATEI = f"{speicher.IMPORT_ORDNER}/zkl.json"
QUELLE = "laut ZKL"
NICHT_BELEGBAR = ("lager", "büro", "buero", "appartment", "apartment", "musterzimmer", "pflegebett")
_ZIMMER = re.compile(r"^\s*([\d.]+)-([A-Z]?\d+)(.*)$")
SPALTEN = {                       # Feld → Anfang der Spaltenüberschrift (ohne Zeilenumbrüche, klein)
    "zimmer": "zimmer nummer", "name": "name", "von": "von", "bis": "bis", "kundengruppe": "kunden- gruppe",
    "tn_id": "id rios", "anmerkung": "interne anmerkungen", "zimmer_daten": "zimmer daten", "reno": "reno",
    "boden": "art", "kuerzel": "x1",
}


@dataclass
class ZklZimmer:
    zimmer: str                   # Zimmer-ID (GS-…)
    roh: str                      # Zimmernummer wie in der ZKL (z. B. „3.3-112 DZ“)
    kundengruppe: str = ""
    tn_id: str = ""
    name: str = ""
    von: date | None = None
    bis: date | None = None
    anmerkung: str = ""
    zimmer_daten: str = ""
    reno: str = ""
    boden: str = ""
    kuerzel: str = ""

    @property
    def tier(self) -> bool:
        return "tier" in self.kundengruppe.lower()

    @property
    def gaestehaus(self) -> bool:
        return "SGH" in (self.kundengruppe.upper(), self.kuerzel.upper())

    @property
    def tier_pool(self) -> bool:
        return "tier" in self.zimmer_daten.lower()

    @property
    def nicht_belegbar(self) -> str:
        """Grund, wenn die ZKL das Zimmer als Lager/Büro/Appartment o. Ä. führt (sonst "")."""
        for text in (self.zimmer_daten, self.roh):
            if any(w in text.lower() for w in NICHT_BELEGBAR) and "ehem" not in text.lower():   # „ehem. Appartment“
                return text.strip()
        return ""

    @property
    def info(self) -> str:
        """Kurztext für die Zimmerkarte."""
        teile = [f"Boden {self.boden}" if self.boden else "", f"Reno {self.reno}" if self.reno else "",
                 self.zimmer_daten, self.kundengruppe and f"Kundengruppe {self.kundengruppe}"]
        return " · ".join(t for t in teile if t)

    def zu_dict(self) -> dict:
        d = asdict(self)
        d["von"] = self.von.isoformat() if self.von else None
        d["bis"] = self.bis.isoformat() if self.bis else None
        return d

    @staticmethod
    def aus_dict(d: dict) -> "ZklZimmer":
        felder = {k: v for k, v in d.items() if k in ZklZimmer.__dataclass_fields__}
        for k in ("von", "bis"):
            felder[k] = date.fromisoformat(felder[k]) if felder.get(k) else None
        return ZklZimmer(**felder)


@dataclass
class Zkl:
    datei: str
    blatt: str
    importiert: str
    zimmer: dict[str, ZklZimmer] = field(default_factory=dict)
    unbekannt: list[str] = field(default_factory=list)       # Zimmernummern, die es im Zimmerplan nicht gibt
    uebernommen: list[dict] = field(default_factory=list)    # als belegt übernommene ZKL-Bewohner
    geprueft: list[str] = field(default_factory=list)        # Abgleich-Punkte, die als geprüft abgelegt sind

    @property
    def stand(self) -> date | None:
        """Neuester Anreisetag in der Liste – ungefährer Stand."""
        tage = [z.von for z in self.zimmer.values() if z.von]
        return max(tage) if tage else None


# ---------------------------------------------------------------------------------------
# Lesen
# ---------------------------------------------------------------------------------------

def _norm(text) -> str:
    return re.sub(r"\s+", " ", str(text or "").replace("\n", " ")).strip().lower()


def _datum(wert) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, str):
        for f in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(wert.strip(), f).date()
            except ValueError:
                pass
    return None


def _text(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float) and wert.is_integer():
        wert = int(wert)
    return re.sub(r"\s+", " ", str(wert)).strip()


def zimmer_id(roh: str, bekannt: set[str]) -> str | None:
    """„3.3-112 DZ“ → GS-3.2-112, wenn es die Nummer nur in Haus 3.2 gibt (die ZKL nennt beide „3.3“)."""
    m = _ZIMMER.match(roh or "")
    if not m:
        return None
    haus, nr = m.group(1), m.group(2)
    zid = f"GS-{haus}-{nr}"
    if zid not in bekannt and haus in ("3.2", "3.3"):
        anders = f"GS-{'3.2' if haus == '3.3' else '3.3'}-{nr}"
        if anders in bekannt:
            return anders
    return zid


def _blatt_lesen(zeilen: list[tuple], bekannt: set[str]) -> tuple[dict[str, ZklZimmer], list[str]] | None:
    kopf = None
    for i, z in enumerate(zeilen[:15]):
        namen = [_norm(c) for c in z]
        if any(n.startswith("zimmer nummer") for n in namen) and "name" in namen:
            kopf = i, namen
            break
    if kopf is None:
        return None
    i, namen = kopf
    spalte = {}
    for feld, anfang in SPALTEN.items():
        treffer = [j for j, n in enumerate(namen) if (n == anfang if feld in ("name", "von", "bis", "boden", "reno", "kuerzel")
                                                      else n.startswith(anfang))]
        if treffer:
            spalte[feld] = treffer[0]
    if not {"zimmer", "name"} <= set(spalte):
        return None
    erg: dict[str, ZklZimmer] = {}
    unbekannt = []

    def wert(z, feld):
        j = spalte.get(feld)
        return z[j] if j is not None and j < len(z) else None

    for z in zeilen[i + 1:]:
        roh = _text(wert(z, "zimmer"))
        zid = zimmer_id(roh, bekannt)
        if not zid:
            continue
        if zid not in bekannt:
            unbekannt.append(roh)
        tn = wert(z, "tn_id")
        erg[zid] = ZklZimmer(zid, roh, _text(wert(z, "kundengruppe")), _text(tn) if tn not in (None, 0) else "",
                             _text(wert(z, "name")), _datum(wert(z, "von")), _datum(wert(z, "bis")),
                             _text(wert(z, "anmerkung")), _text(wert(z, "zimmer_daten")), _text(wert(z, "reno")),
                             _text(wert(z, "boden")), _text(wert(z, "kuerzel")))
    return erg, unbekannt


def lesen(pfad: Path, bekannt: set[str] | None = None) -> Zkl:
    """Liest alle Blätter mit Zimmerliste und nimmt das mit den neuesten Daten."""
    import warnings

    import openpyxl

    if bekannt is None:
        bekannt = {z.id for z in zp.laden().zimmer}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")              # z. B. ungültiger Druckbereich – für die Daten egal
        wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    beste = None
    try:
        for ws in wb.worksheets:
            if (ws.max_column or 0) > 500 or (ws.max_row or 0) > 5000:
                continue                       # leere Riesenblätter (formatierte ganze Spalten) überspringen
            zeilen = [tuple(r) for r in ws.iter_rows(values_only=True)]
            gelesen = _blatt_lesen(zeilen, bekannt)
            if not gelesen or not gelesen[0]:
                continue
            neueste = max((z.von for z in gelesen[0].values() if z.von), default=date.min)
            if beste is None or (neueste, len(gelesen[0])) > beste[0]:
                beste = ((neueste, len(gelesen[0])), ws.title, gelesen)
    finally:
        wb.close()
    if beste is None:
        raise ValueError("Keine Zimmerliste gefunden (Spalten „Zimmer Nummer“ und „Name“ fehlen).")
    _r, blatt, (zimmer, unbekannt) = beste
    alt = laden()
    return Zkl(Path(pfad).name, blatt, datetime.now().isoformat(timespec="seconds"), zimmer, unbekannt,
               alt.uebernommen if alt else [], alt.geprueft if alt else [])


def laden() -> Zkl | None:
    roh = speicher.lesen(DATEI, {})
    if not roh.get("zimmer"):
        return None
    return Zkl(roh.get("datei", ""), roh.get("blatt", ""), roh.get("importiert", ""),
               {k: ZklZimmer.aus_dict(v) for k, v in roh["zimmer"].items()}, roh.get("unbekannt", []),
               roh.get("uebernommen", []), roh.get("geprueft", []))


def speichern(zkl: Zkl) -> None:
    speicher.schreiben(DATEI, {"datei": zkl.datei, "blatt": zkl.blatt, "importiert": zkl.importiert,
                               "zimmer": {k: v.zu_dict() for k, v in zkl.zimmer.items()}, "unbekannt": zkl.unbekannt,
                               "uebernommen": zkl.uebernommen, "geprueft": zkl.geprueft})


# ---------------------------------------------------------------------------------------
# In der Lage nutzen
# ---------------------------------------------------------------------------------------

def _gleich(a: str, b: str) -> bool:
    wa, wb = zp._namenswoerter(a), zp._namenswoerter(b)
    return bool(wa and wb) and (wa <= wb or wb <= wa)


def anwenden(lage: zp.Lage, zkl: Zkl | None) -> None:
    """ZKL an die Lage hängen (Tier, TN-ID, Zimmer-Infos) und übernommene Bewohner als belegt eintragen –
    außer der Gebäudeplan kennt die Person inzwischen (dann gilt er)."""
    lage.zkl = zkl.zimmer if zkl else {}
    if zkl is None:
        return
    for u in zkl.uebernommen:
        zid, von, bis = u["zimmer"], date.fromisoformat(u["von"]), date.fromisoformat(u["bis"])
        if zid not in lage.belegungen:
            continue
        bekannt = any(b.art == "belegt" and b.name and _gleich(b.name, u["name"]) and b.von <= bis and von <= b.bis
                      for bs in lage.belegungen.values() for b in bs)
        if not bekannt:
            lage.belegungen[zid].append(Belegung(zid, von, bis, "belegt", u["name"],
                                                 f"{u.get('kundengruppe') or 'Reha'} · {QUELLE}", QUELLE))


def eintrag(lage: zp.Lage, b: Belegung) -> ZklZimmer | None:
    """ZKL-Zeile zu einer Belegung (gleiches Zimmer, gleicher Name)."""
    z = getattr(lage, "zkl", {}).get(b.zimmer)
    return z if z is not None and z.name and b.name and _gleich(z.name, b.name) else None


def hat_tier(lage: zp.Lage, b: Belegung) -> bool | None:
    """Tier laut ZKL (None = unbekannt)."""
    z = eintrag(lage, b)
    return z.tier if z is not None else None


def tn_id(lage: zp.Lage, b: Belegung) -> str:
    z = eintrag(lage, b)
    return z.tn_id if z is not None else ""


# ---------------------------------------------------------------------------------------
# Abgleich
# ---------------------------------------------------------------------------------------

@dataclass
class Punkt:
    art: str                      # bewohnt | nicht_belegbar | gaeste | tiere | anderes_zimmer | unbekannt
    zimmer: str
    text: str
    aktion: str = ""              # Beschriftung der Aktion ("" = nur Info)
    daten: dict = field(default_factory=dict)

    @property
    def schluessel(self) -> str:
        return f"{self.art}|{self.zimmer}|{self.daten.get('name', '')}"


VORAUSWAHL = {"bewohnt", "nicht_belegbar", "tiere"}   # ZKL ist ab jetzt statisch: diese einmalig übernehmen
ARTEN = {
    "bewohnt": "Laut ZKL bewohnt, im Gebäudeplan frei",
    "nicht_belegbar": "Laut ZKL nicht belegbar (Lager, Büro, Appartment …)",
    "gaeste": "Laut ZKL Gästehaus (SGH)",
    "tiere": "Tier-Zimmer laut ZKL anders als im Zimmerplan",
    "anderes_zimmer": "Person wohnt inzwischen in einem anderen Zimmer (Gebäudeplan gilt)",
    "unbekannt": "Zimmer aus der ZKL, die es im Zimmerplan nicht gibt",
}


def abgleich(stand: zp.Planstand, zkl: Zkl, heute: date | None = None, alle: bool = False) -> list[Punkt]:
    """Abweichungen ZKL ↔ Zimmerplan. Als geprüft abgelegte Punkte nur mit ``alle``."""
    heute = heute or date.today()
    lage, stamm = stand.lage, {z.id: z for z in stand.zimmer}
    personen = [(zid, b) for zid, bs in lage.belegungen.items() for b in bs
                if b.art in ("belegt", "geplant") and b.name and b.name != zp.PIVOT_NAME]
    erg: list[Punkt] = []
    uebernommen = {(u["zimmer"], u["name"]) for u in zkl.uebernommen}
    for zid, e in sorted(zkl.zimmer.items(), key=lambda x: x[0]):
        z = stamm.get(zid)
        if z is None:
            continue
        k = zid.removeprefix("GS-")
        if e.name and e.von and (e.bis is None or e.bis >= heute) and (zid, e.name) not in uebernommen:
            woanders = [(x, b) for x, b in personen if _gleich(b.name, e.name) and b.bis >= heute]
            if woanders and all(x != zid for x, _b in woanders):
                erg.append(Punkt("anderes_zimmer", zid, f"{e.name}: laut ZKL {k}, laut Plan "
                                 f"{woanders[0][0].removeprefix('GS-')}", daten={"name": e.name}))
            elif not woanders and z.aktiv and lage.freie_betten(zid, max(heute, e.von), e.bis or max(heute, e.von)) > 0:
                bis = e.bis or date(heute.year + 1, heute.month, 1)
                erg.append(Punkt("bewohnt", zid, f"{k}: laut ZKL {e.name}" + (f" (TN {e.tn_id})" if e.tn_id else "")
                                 + f", {e.kundengruppe or 'ohne Kundengruppe'}, {e.von:%d.%m.%Y} – "
                                 + (f"{e.bis:%d.%m.%Y}" if e.bis else "offen") + " – im Gebäudeplan frei. "
                                 "Ausgezogen oder fehlt im Gebäudeplan?", "als belegt übernehmen",
                                 {"name": e.name, "von": e.von.isoformat(),
                                  "bis": bis.isoformat(), "kundengruppe": e.kundengruppe}))
        if (grund := e.nicht_belegbar) and z.aktiv:
            erg.append(Punkt("nicht_belegbar", zid, f"{k}: laut ZKL „{grund}“ – im Zimmerplan belegbar"
                             + (" (Gästezimmer)" if z.gaeste else ""), "gehört nicht zum Internat"))
        if e.gaestehaus and not z.gaeste and z.aktiv:
            erg.append(Punkt("gaeste", zid, f"{k}: laut ZKL Gästehaus (SGH) – im Zimmerplan normales Zimmer",
                             "als Gästezimmer markieren"))
        if e.tier_pool != z.tiere and z.aktiv and z.haus == "3.1":
            erg.append(Punkt("tiere", zid, f"{k}: laut ZKL {'Tier-Pool' if e.tier_pool else 'kein Tier-Zimmer'} – im "
                             f"Zimmerplan {'Tier-Zimmer' if z.tiere else 'kein Tier-Zimmer'}",
                             "Tier-Zimmer" if e.tier_pool else "kein Tier-Zimmer", {"tiere": e.tier_pool}))
    if zkl.unbekannt:
        erg.append(Punkt("unbekannt", "", ", ".join(zkl.unbekannt[:40]) + (" …" if len(zkl.unbekannt) > 40 else "")))
    return [p for p in erg if alle or p.schluessel not in zkl.geprueft]


def umsetzen(stand: zp.Planstand, zkl: Zkl, gewaehlt: list[Punkt], geprueft: list[Punkt]) -> dict[str, int]:
    """Gewählte Punkte übernehmen, die übrigen als geprüft ablegen. Gibt Anzahl je Art zurück."""
    zimmer = {z.id: z for z in stand.zimmer}
    zaehler: dict[str, int] = {}
    stamm_geaendert = False
    for p in gewaehlt:
        z = zimmer.get(p.zimmer)
        if p.art == "bewohnt":
            zkl.uebernommen.append({"zimmer": p.zimmer, **p.daten})
        elif z is None:
            continue
        elif p.art == "nicht_belegbar":
            z.aktiv, stamm_geaendert = False, True
        elif p.art == "gaeste":
            z.gaeste, stamm_geaendert = True, True
        elif p.art == "tiere":
            z.tiere, stamm_geaendert = bool(p.daten.get("tiere")), True
        else:
            continue
        zaehler[p.art] = zaehler.get(p.art, 0) + 1
    zkl.geprueft = sorted(set(zkl.geprueft) | {p.schluessel for p in geprueft} | {p.schluessel for p in gewaehlt})
    if stamm_geaendert:
        zp.stammdaten_speichern(stand.zimmer, stand.plan)
    speichern(zkl)
    return zaehler
