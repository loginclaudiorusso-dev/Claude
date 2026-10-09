"""Inventar: Gegenstände (mit Bestand) und ihre Verteilung auf die Zimmer.

Gespeichert in ``inventar.json`` neben den anderen Daten:
    {"gegenstaende": [{id, name, kategorie, bestand, notiz}], "zimmer": {Zimmer-ID: {Gegenstand-ID: Anzahl}}}

``bestand`` ist die Zahl, die insgesamt vorhanden ist (None = nicht gezählt); was nicht in Zimmern
steht, gilt als „im Lager“. Rein, ohne Qt.
"""

from __future__ import annotations

import re
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path

from . import speicher

DATEI = "inventar.json"
KATEGORIEN = ["Möbel", "Bett & Wäsche", "Elektro", "Bad", "Küche", "Reinigung", "Sonstiges"]


@dataclass
class Gegenstand:
    name: str
    kategorie: str = "Sonstiges"
    bestand: int | None = None      # insgesamt vorhanden (None = nicht gezählt)
    notiz: str = ""
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:10])


@dataclass
class Inventar:
    gegenstaende: list[Gegenstand] = field(default_factory=list)
    zimmer: dict[str, dict[str, int]] = field(default_factory=dict)

    # ---- Gegenstände ---------------------------------------------------------------------

    def gegenstand(self, gid: str) -> Gegenstand | None:
        return next((g for g in self.gegenstaende if g.id == gid), None)

    def nach_name(self, name: str) -> Gegenstand | None:
        n = name.strip().lower()
        return next((g for g in self.gegenstaende if g.name.strip().lower() == n), None)

    def gegenstand_speichern(self, g: Gegenstand) -> None:
        """Neu oder geändert. Namen sind eindeutig (Groß-/Kleinschreibung egal)."""
        if not g.name.strip():
            raise ValueError("Bitte einen Namen eingeben.")
        doppelt = self.nach_name(g.name)
        if doppelt is not None and doppelt.id != g.id:
            raise ValueError(f"„{doppelt.name}“ gibt es schon.")
        g.name = g.name.strip()
        alt = self.gegenstand(g.id)
        if alt is None:
            self.gegenstaende.append(g)
        else:
            self.gegenstaende[self.gegenstaende.index(alt)] = g
        self.gegenstaende.sort(key=lambda x: (KATEGORIEN.index(x.kategorie) if x.kategorie in KATEGORIEN else 99,
                                              x.name.lower()))

    def gegenstand_loeschen(self, gid: str) -> None:
        self.gegenstaende = [g for g in self.gegenstaende if g.id != gid]
        for ausstattung in self.zimmer.values():
            ausstattung.pop(gid, None)
        self._aufraeumen()

    # ---- Zuordnung -----------------------------------------------------------------------

    def anzahl(self, zid: str, gid: str) -> int:
        return self.zimmer.get(zid, {}).get(gid, 0)

    def setzen(self, zid: str, gid: str, anzahl: int) -> None:
        if anzahl <= 0:
            self.zimmer.get(zid, {}).pop(gid, None)
        else:
            self.zimmer.setdefault(zid, {})[gid] = int(anzahl)
        self._aufraeumen()

    def hinzufuegen(self, zimmer: list[tuple[str, int]], gid: str, anzahl: int, je_bett: bool = False) -> None:
        """``zimmer``: (Zimmer-ID, Betten). Erhöht die Anzahl in jedem Zimmer (je Bett: × Betten)."""
        for zid, betten in zimmer:
            self.setzen(zid, gid, self.anzahl(zid, gid) + anzahl * (max(betten, 1) if je_bett else 1))

    def entfernen(self, zids: list[str], gid: str) -> None:
        for zid in zids:
            self.setzen(zid, gid, 0)

    def kopieren(self, von: str, nach: list[str]) -> None:
        """Ausstattung von Zimmer ``von`` in die Zimmer ``nach`` übernehmen (ersetzt deren Ausstattung)."""
        vorlage = dict(self.zimmer.get(von, {}))
        for zid in nach:
            if zid != von:
                self.zimmer[zid] = dict(vorlage)
        self._aufraeumen()

    def _aufraeumen(self) -> None:
        self.zimmer = {z: a for z, a in self.zimmer.items() if a}

    # ---- Auswertung ----------------------------------------------------------------------

    def im_zimmer(self, zid: str) -> list[tuple[Gegenstand, int]]:
        a = self.zimmer.get(zid, {})
        return [(g, a[g.id]) for g in self.gegenstaende if g.id in a]

    def wo(self, gid: str) -> list[tuple[str, int]]:
        return sorted((zid, a[gid]) for zid, a in self.zimmer.items() if gid in a)

    def verteilt(self, gid: str) -> int:
        return sum(a.get(gid, 0) for a in self.zimmer.values())

    def frei(self, gid: str) -> int | None:
        """Im Lager (Bestand minus verteilt); None, wenn der Bestand nicht gezählt wird. Negativ = zu viel verteilt."""
        g = self.gegenstand(gid)
        return None if g is None or g.bestand is None else g.bestand - self.verteilt(gid)

    def kurztext(self, zid: str) -> str:
        """„2× Bett, 2× Schrank, Schreibtisch“."""
        return ", ".join(f"{n}× {g.name}" if n > 1 else g.name for g, n in self.im_zimmer(zid))

    def je_haus(self, zimmer_haus: dict[str, str]) -> dict[str, dict[str, int]]:
        """Gegenstand-ID → {Haus: Anzahl}."""
        erg: dict[str, dict[str, int]] = {}
        for zid, a in self.zimmer.items():
            haus = zimmer_haus.get(zid, "?")
            for gid, n in a.items():
                erg.setdefault(gid, {})[haus] = erg.setdefault(gid, {}).get(haus, 0) + n
        return erg

    def zu_dict(self) -> dict:
        return {"gegenstaende": [asdict(g) for g in self.gegenstaende], "zimmer": self.zimmer}

    @staticmethod
    def aus_dict(d: dict) -> "Inventar":
        gs = []
        for x in d.get("gegenstaende", []):
            try:
                gs.append(Gegenstand(x["name"], x.get("kategorie") or "Sonstiges", x.get("bestand"), x.get("notiz") or "",
                                     x.get("id") or uuid.uuid4().hex[:10]))
            except (KeyError, TypeError):
                continue
        bekannt = {g.id for g in gs}
        zimmer = {zid: {gid: int(n) for gid, n in a.items() if gid in bekannt and int(n) > 0}
                  for zid, a in (d.get("zimmer") or {}).items()}
        return Inventar(gs, {z: a for z, a in zimmer.items() if a})


def laden() -> Inventar:
    return Inventar.aus_dict(speicher.lesen(DATEI, {}))


def speichern(inv: Inventar) -> None:
    speicher.schreiben(DATEI, inv.zu_dict())


# ---------------------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------------------

def _sortier(zid: str) -> tuple:
    """Nach Haus, dann numerisch nach Zimmernummer (E01 vor 101)."""
    teile = zid.removeprefix("GS-").split("-")
    haus = float(teile[0]) if re.fullmatch(r"[\d.]+", teile[0]) and teile[0].count(".") <= 1 else 99.0
    m = re.search(r"(\d+)$", teile[-1])
    return haus, int(m.group(1)) if m else 0, teile[-1]


def excel_speichern(pfad: Path, inv: Inventar, zimmer: list) -> None:
    """Drei Blätter: Zimmer × Gegenstand (Matrix), Gegenstände mit Bestand, Liste je Zimmer."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fett, kopf_farbe = Font(bold=True), PatternFill("solid", fgColor="DDEBE8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Zimmer"
    gs = inv.gegenstaende
    ws.append(["Zimmer", "Haus", "Etage", "Betten"] + [g.name for g in gs])
    for z in sorted(zimmer, key=lambda z: _sortier(z.id)):
        if not z.aktiv:
            continue
        ws.append([z.kurz, z.haus, z.etage, z.betten] + [inv.anzahl(z.id, g.id) or None for g in gs])
    ws.append([])
    ws.append(["Summe", "", "", ""] + [inv.verteilt(g.id) for g in gs])
    for c in ws[1]:
        c.font, c.fill = fett, kopf_farbe
        c.alignment = Alignment(text_rotation=0, wrap_text=True, vertical="top")
    for c in ws[ws.max_row]:
        c.font = fett
    ws.freeze_panes = "E2"
    ws.column_dimensions["A"].width = 11
    for i in range(len(gs)):
        ws.column_dimensions[ws.cell(1, 5 + i).column_letter].width = max(8, min(18, len(gs[i].name) + 2))

    ws2 = wb.create_sheet("Gegenstände")
    ws2.append(["Gegenstand", "Kategorie", "Bestand", "in Zimmern", "im Lager", "Zimmer", "Notiz"])
    for g in gs:
        frei = inv.frei(g.id)
        ws2.append([g.name, g.kategorie, g.bestand, inv.verteilt(g.id), frei, len(inv.wo(g.id)), g.notiz])
    for c in ws2[1]:
        c.font, c.fill = fett, kopf_farbe
    for col, w in zip("ABCDEFG", (26, 14, 10, 11, 10, 9, 40)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Liste je Zimmer")
    ws3.append(["Zimmer", "Gegenstand", "Kategorie", "Anzahl"])
    for zid in sorted(inv.zimmer, key=_sortier):
        for g, n in inv.im_zimmer(zid):
            ws3.append([zid.removeprefix("GS-"), g.name, g.kategorie, n])
    for c in ws3[1]:
        c.font, c.fill = fett, kopf_farbe
    for col, w in zip("ABCD", (11, 26, 14, 8)):
        ws3.column_dimensions[col].width = w
    wb.save(pfad)


def druck_html(inv: Inventar, zimmer: list, nur_haus: str | None = None) -> str:
    """Inventarliste je Zimmer zum Abhaken beim Rundgang (mehrere Zimmer je Seite)."""
    from html import escape

    teile = ["<style>body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 9.5pt; color: #1f2a2a; }"
             "h1 { font-size: 15pt; margin: 0 0 6px 0; } h2 { font-size: 11pt; margin: 14px 0 2px 0; color: #0f5e57; }"
             ".klein { color: #5b6b6b; font-size: 8.5pt; } td { padding: 2px 8px 2px 0; }</style>",
             "<h1>Inventarliste" + (f" Haus {escape(nur_haus)}" if nur_haus else "") + "</h1>",
             "<div class='klein'>Je Zimmer abhaken: ☐ vorhanden · Abweichungen notieren</div>"]
    for z in sorted(zimmer, key=lambda z: _sortier(z.id)):
        if not z.aktiv or (nur_haus and z.haus != nur_haus):
            continue
        liste = inv.im_zimmer(z.id)
        if not liste:
            continue
        teile.append(f"<h2>{escape(z.kurz)} <span class='klein'>· {z.betten} Bett{'en' if z.betten > 1 else ''}"
                     "</span></h2><table cellspacing='0'>")
        for g, n in liste:
            teile.append(f"<tr><td>☐</td><td align='right'><b>{n}×</b></td><td>{escape(g.name)}</td>"
                         f"<td class='klein'>{escape(g.kategorie)}</td><td class='klein'>________________</td></tr>")
        teile.append("</table>")
    return "".join(teile)
