"""Mängel je Zimmer: Was im Zimmer kaputt oder schmutzig ist – zum Filtern und für das Facility-Management (FM).

In der Zimmerkarte stehen Vorschläge (Teppich beschädigt, Schimmel, Heizung …) als Chips: antippen = Mangel offen
(seit heute), noch einmal = behoben. Behobene Mängel bleiben mit Datum in der Datei (Verlauf je Zimmer); gezeigt und
gefiltert werden nur offene. Rein, ohne Qt. Gespeichert in ``zimmer_maengel.json``.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from html import escape

from . import protokoll, speicher

DATEI = "zimmer_maengel.json"

KATALOG: list[tuple[str, str]] = [
    ("boden_kaputt", "Teppich/Boden beschädigt"),
    ("boden_dreckig", "Teppich/Boden verschmutzt"),
    ("wand", "Wand/Tapete beschädigt"),
    ("schimmel", "Schimmel/Feuchtigkeit"),
    ("fenster", "Fenster/Rollo defekt"),
    ("heizung", "Heizung defekt"),
    ("bad", "Bad/Dusche/WC defekt"),
    ("abfluss", "Abfluss verstopft"),
    ("wasser", "Wasserhahn tropft/undicht"),
    ("elektrik", "Licht/Steckdose defekt"),
    ("moebel", "Möbel beschädigt"),
    ("matratze", "Matratze erneuern"),
    ("tuer", "Tür/Schloss/Schlüssel defekt"),
    ("rauchmelder", "Rauchmelder defekt/fehlt"),
    ("geruch", "Geruch (Rauch u. a.)"),
    ("ungeziefer", "Ungeziefer"),
    ("grundreinigung", "Grundreinigung nötig"),
    ("renovierung", "Renovierung nötig"),        # z. B. ZKL „-> Reno“, „Reno1/Reno2“
    ("sonstiges", "Sonstiges"),
]
NAMEN = dict(KATALOG)
# Diese Arten machen ein Zimmer von sich aus unbewohnbar; bei allen anderen entscheidet der Nutzer je Mangel
# („macht unbewohnbar“ in der Zimmerkarte) – ein kaputtes Rollo ist kein Grund, eine fehlende Matratze schon.
UNBEWOHNBAR_STANDARD = {"schimmel", "ungeziefer"}


@dataclass
class Mangel:
    zimmer: str
    art: str
    seit: date
    behoben: date | None = None
    notiz: str = ""
    sperrt: bool | None = None      # None = Standard der Art (UNBEWOHNBAR_STANDARD)

    @property
    def unbewohnbar(self) -> bool:
        return self.sperrt if self.sperrt is not None else self.art in UNBEWOHNBAR_STANDARD

    @property
    def text(self) -> str:
        return NAMEN.get(self.art, self.art) + (f" ({self.notiz})" if self.notiz else "")

    def zu_dict(self) -> dict:
        d = {"zimmer": self.zimmer, "art": self.art, "seit": self.seit.isoformat(),
             "behoben": self.behoben.isoformat() if self.behoben else None, "notiz": self.notiz}
        if self.sperrt is not None:
            d["sperrt"] = self.sperrt
        return d

    @staticmethod
    def aus_dict(d: dict) -> "Mangel":
        return Mangel(d["zimmer"], d["art"], date.fromisoformat(d["seit"]),
                      date.fromisoformat(d["behoben"]) if d.get("behoben") else None, d.get("notiz", ""),
                      d["sperrt"] if isinstance(d.get("sperrt"), bool) else None)


def laden() -> list[Mangel]:
    erg = []
    for d in speicher.lesen(DATEI, {}).get("maengel", []):
        try:
            erg.append(Mangel.aus_dict(d))
        except (KeyError, ValueError, TypeError):
            continue
    return erg


def offen(zimmer: str | None = None) -> list[Mangel]:
    return [m for m in laden() if m.behoben is None and (zimmer is None or m.zimmer == zimmer)]


def offen_je_zimmer() -> dict[str, list[Mangel]]:
    erg: dict[str, list[Mangel]] = {}
    for m in offen():
        erg.setdefault(m.zimmer, []).append(m)
    return erg


def unbewohnbar_je_zimmer() -> dict[str, str]:
    """Zimmer-ID → Grund, für alle Zimmer mit einem offenen Mangel, der sie unbewohnbar macht."""
    erg: dict[str, list[str]] = {}
    for m in offen():
        if m.unbewohnbar:
            erg.setdefault(m.zimmer, []).append(m.text)
    return {z: "; ".join(t) for z, t in erg.items()}


def sperren_setzen(zimmer: str, art: str, unbewohnbar: bool) -> bool:
    """Offenen Mangel als „macht unbewohnbar“ markieren (oder nicht). True, wenn sich etwas geändert hat."""
    alle = laden()
    treffer = [m for m in alle if m.zimmer == zimmer and m.art == art and m.behoben is None]
    if not treffer or all(m.unbewohnbar == unbewohnbar for m in treffer):
        return False
    for m in treffer:
        m.sperrt = unbewohnbar
    speicher.schreiben(DATEI, {"maengel": [m.zu_dict() for m in alle]})
    protokoll.eintragen([{"aktion": "Zimmer unbewohnbar" if unbewohnbar else "Zimmer wieder bewohnbar",
                          "name": NAMEN.get(art, art), "person": "", "alt": "", "neu": zimmer,
                          "von": date.today().isoformat(), "quelle": "Zimmerkarte"}])
    return True


def setzen(zimmer: str, art: str, an: bool, heute: date | None = None, notiz: str = "") -> bool:
    """Mangel öffnen (an) bzw. als behoben markieren (aus). True, wenn sich etwas geändert hat."""
    heute = heute or date.today()
    alle = laden()
    offene = [m for m in alle if m.zimmer == zimmer and m.art == art and m.behoben is None]
    if an and not offene:
        alle.append(Mangel(zimmer, art, heute, None, notiz.strip()))
    elif not an and offene:
        for m in offene:
            m.behoben = heute
    else:
        return False
    speicher.schreiben(DATEI, {"maengel": [m.zu_dict() for m in alle]})
    protokoll.eintragen([{"aktion": "Mangel gemeldet" if an else "Mangel behoben", "name": NAMEN.get(art, art),
                          "person": "", "alt": zimmer if not an else "", "neu": zimmer if an else "",
                          "von": heute.isoformat(), "quelle": "Zimmerkarte"}])
    return True


def fm_text(offene: list[Mangel], kurz=lambda zid: zid.removeprefix("GS-")) -> str:
    """Zum Einfügen in eine Mail an das Facility-Management – je Zimmer eine Zeile."""
    je: dict[str, list[Mangel]] = {}
    for m in offene:
        je.setdefault(m.zimmer, []).append(m)
    zeilen = [f"Offene Mängel – Internat Goslar, Stand {datetime.now():%d.%m.%Y}", ""]
    for zid in sorted(je, key=_sort):
        zeilen.append(f"Zimmer {kurz(zid)}: " + "; ".join(f"{m.text} (seit {m.seit:%d.%m.})" for m in je[zid]))
    return "\n".join(zeilen)


@dataclass
class BerichtZeile:
    zimmer: str            # Zimmer-ID
    kurz: str              # „2-402“
    haus: str
    art: str
    mangel: str            # Name aus dem Katalog
    details: str           # Notiz (z. B. „ZKL: WC-Deckel lose“)
    seit: date
    tage: int              # wie lange schon offen
    zimmer_heute: str = "" # frei | belegt | … – damit das FM weiß, ob es klopfen muss (keine Namen)
    dringend: bool = False # macht das Zimmer unbewohnbar → zuerst
    @property
    def schluessel(self) -> str:
        return f"{self.zimmer}|{self.art}"


ZIMMER_HEUTE = {"frei": "frei", "teilweise": "1 Bett frei", "belegt": "belegt", "geplant": "frei (geplant)",
                "gesperrt": "gesperrt"}


def bericht(offene: list[Mangel], kurz=lambda zid: zid.removeprefix("GS-"), status: dict[str, str] | None = None,
            heute: date | None = None) -> list[BerichtZeile]:
    """Eine Zeile je offenem Mangel, nach Haus und Zimmer sortiert."""
    heute = heute or date.today()
    status = status or {}
    erg = [BerichtZeile(m.zimmer, kurz(m.zimmer), m.zimmer.removeprefix("GS-").partition("-")[0], m.art,
                        NAMEN.get(m.art, m.art), m.notiz, m.seit, (heute - m.seit).days,
                        ZIMMER_HEUTE.get(status.get(m.zimmer, ""), status.get(m.zimmer, "")), m.unbewohnbar)
           for m in offene]
    return sorted(erg, key=lambda r: (_sort(r.zimmer), not r.dringend, r.mangel))


def fm_html(offene: list[Mangel], kurz=lambda zid: zid.removeprefix("GS-"), status: dict[str, str] | None = None,
            heute: date | None = None) -> str:
    """Mängelbericht für das FM: Überblick (je Art, je Haus), dann je Haus eine Tabelle zum Abhaken."""
    from .drucklisten import _STIL

    heute = heute or date.today()
    zeilen = bericht(offene, kurz, status, heute)
    teile = [_STIL, "<h1>Mängelbericht für das Facility-Management</h1>",
             f"<div class='klein'>Internat Goslar · {len(zeilen)} offene Mängel in {len({r.zimmer for r in zeilen})} "
             f"Zimmern · Stand {datetime.now():%d.%m.%Y %H:%M} · erledigte bitte mit Datum abzeichnen</div>"]
    if not zeilen:
        return "".join(teile + ["<p class='leer'>Keine offenen Mängel.</p>"])
    dringend = [r for r in zeilen if r.dringend]
    if dringend:                                   # unbewohnbare Zimmer zuerst – sie fehlen bei der Belegung
        teile.append(f"<h2>Dringend – Zimmer nicht bewohnbar ({len({r.zimmer for r in dringend})})</h2>"
                     "<table class='liste' width='100%' cellspacing='0'><tr><th width='9%'>Zimmer</th>"
                     "<th width='10%'>heute</th><th width='20%'>Mangel</th><th>Details</th><th width='8%'>seit</th>"
                     "<th width='11%'>erledigt am</th><th width='9%'>Kürzel</th></tr>")
        teile += [f"<tr><td><b>{escape(r.kurz)}</b></td><td>{escape(r.zimmer_heute)}</td><td><b>{escape(r.mangel)}</b></td>"
                  f"<td>{escape(r.details)}</td><td>{r.seit:%d.%m.%y}</td><td></td><td></td></tr>" for r in dringend]
        teile.append("</table>")
    je_art: dict[str, int] = {}
    je_haus: dict[str, int] = {}
    for r in zeilen:
        je_art[r.mangel] = je_art.get(r.mangel, 0) + 1
        je_haus[r.haus] = je_haus.get(r.haus, 0) + 1
    teile.append("<h2>Überblick</h2><table class='liste' cellspacing='0'><tr><th>Mangel</th><th>Anzahl</th></tr>")
    teile += [f"<tr><td>{escape(a)}</td><td>{n}</td></tr>" for a, n in sorted(je_art.items(), key=lambda x: -x[1])]
    teile.append("</table><div class='klein'>Je Haus: "
                 + " · ".join(f"Haus {escape(h)}: {n}" for h, n in je_haus.items()) + "</div>")
    for haus in je_haus:
        teile.append(f"<h2>Haus {escape(haus)}</h2><table class='liste' width='100%' cellspacing='0'><tr>"
                     "<th width='9%'>Zimmer</th><th width='10%'>heute</th><th width='20%'>Mangel</th><th>Details</th>"
                     "<th width='8%'>seit</th><th width='11%'>erledigt am</th><th width='9%'>Kürzel</th></tr>")
        vorher = ""
        for r in (r for r in zeilen if r.haus == haus):
            erst = r.zimmer != vorher
            vorher = r.zimmer
            teile.append(f"<tr><td><b>{escape(r.kurz) if erst else ''}</b></td><td>{escape(r.zimmer_heute) if erst else ''}</td>"
                         f"<td>{'<b>! ' if r.dringend else ''}{escape(r.mangel)}{' (unbewohnbar)</b>' if r.dringend else ''}</td><td>{escape(r.details)}</td>"
                         f"<td>{r.seit:%d.%m.%y}</td><td></td><td></td></tr>")
        teile.append("</table>")
    return "".join(teile)


EXCEL_SPALTEN = ["Haus", "Zimmer", "Zimmer heute", "Mangel", "Details", "unbewohnbar", "offen seit", "Tage offen",
                 "erledigt am", "Bemerkung FM", "Schlüssel"]


def fm_excel(pfad, offene: list[Mangel], kurz=lambda zid: zid.removeprefix("GS-"), status: dict[str, str] | None = None,
             heute: date | None = None) -> int:
    """Excel zum Abarbeiten: je Mangel eine Zeile; das FM trägt „erledigt am“ (und eine Bemerkung) ein.
    Die ausgefüllte Datei lässt sich mit ``fm_erledigte_lesen`` wieder einlesen. Gibt die Zahl der Zeilen zurück."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    zeilen = bericht(offene, kurz, status, heute)
    wb = Workbook()
    ws = wb.active
    ws.title = "Mängel"
    ws.append(EXCEL_SPALTEN)
    for r in zeilen:
        ws.append([r.haus, r.kurz, r.zimmer_heute, r.mangel, r.details, "ja – dringend" if r.dringend else "", r.seit,
                   r.tage, None, None, r.schluessel])
    kopf = PatternFill("solid", fgColor="E3EFED")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), kopf
    gelb = PatternFill("solid", fgColor="FFF6D5")      # hier trägt das FM ein
    rot = Font(bold=True, color="B42318")
    for c in ws["F"][1:]:
        if c.value:
            c.font = rot
    for zeile in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in zeile:
            c.fill = gelb
    for c in ws["G"][1:]:
        c.number_format = "DD.MM.YYYY"
    for c in ws["I"][1:]:
        c.number_format = "DD.MM.YYYY"
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True,
                        error="Bitte ein Datum eintragen, z. B. 17.10.2026", showErrorMessage=True)
    ws.add_data_validation(dv)
    if zeilen:
        dv.add(f"I2:I{len(zeilen) + 1}")
    for spalte, breite in zip("ABCDEFGHIJK", (7, 10, 13, 28, 50, 14, 12, 10, 13, 30, 22)):
        ws.column_dimensions[spalte].width = breite
    ws.column_dimensions["K"].hidden = True             # zum Wiedereinlesen – bitte nicht ändern
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(len(zeilen), 1) + 1}"
    uebersicht = wb.create_sheet("Überblick")
    uebersicht.append(["Mangel", "Anzahl"])
    je: dict[str, int] = {}
    for r in zeilen:
        je[r.mangel] = je.get(r.mangel, 0) + 1
    for a, n in sorted(je.items(), key=lambda x: -x[1]):
        uebersicht.append([a, n])
    uebersicht.append([])
    uebersicht.append(["Stand", datetime.now().strftime("%d.%m.%Y %H:%M")])
    uebersicht.column_dimensions["A"].width = 32
    wb.save(pfad)
    return len(zeilen)


def _datum_aus(wert, heute: date) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    text = str(wert or "").strip()
    if not text:
        return None
    if text.lower() in ("x", "ja", "erl", "erledigt", "ok"):
        return heute
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def fm_erledigte_lesen(pfad, heute: date | None = None) -> tuple[list[tuple[str, str, date, str]], list[str]]:
    """Ausgefüllte FM-Excel lesen: [(Zimmer-ID, Art, erledigt am, Bemerkung)] und unlesbare Einträge."""
    from openpyxl import load_workbook

    heute = heute or date.today()
    ws = load_workbook(pfad, data_only=True).worksheets[0]
    zeilen = list(ws.iter_rows(values_only=True))
    if not zeilen:
        return [], []
    kopf = [str(c or "").strip().lower() for c in zeilen[0]]
    try:
        i_s, i_e, i_b = kopf.index("schlüssel"), kopf.index("erledigt am"), kopf.index("bemerkung fm")
    except ValueError:
        return [], ["Spalten „erledigt am“/„Schlüssel“ nicht gefunden – ist das die FM-Liste aus dem Programm?"]
    erg, fehler = [], []
    for z in zeilen[1:]:
        schl, wert = z[i_s] if i_s < len(z) else None, z[i_e] if i_e < len(z) else None
        if not schl or wert in (None, ""):
            continue
        d = _datum_aus(wert, heute)
        zid, _, art = str(schl).partition("|")
        if d is None or not art:
            fehler.append(f"{zid.removeprefix('GS-')}: „{wert}“")
            continue
        erg.append((zid, art, d, str(z[i_b] or "").strip() if i_b < len(z) else ""))
    return erg, fehler


def erledigte_uebernehmen(eintraege: list[tuple[str, str, date, str]]) -> int:
    """Offene Mängel als behoben markieren (mit Datum vom FM, Bemerkung an die Notiz). Gibt die Zahl zurück."""
    alle = laden()
    n = 0
    protokolle = []
    for zid, art, tag, bemerkung in eintraege:
        for m in alle:
            if m.zimmer == zid and m.art == art and m.behoben is None:
                m.behoben = tag
                if bemerkung:
                    m.notiz = f"{m.notiz} · FM: {bemerkung}" if m.notiz else f"FM: {bemerkung}"
                n += 1
                protokolle.append({"aktion": "Mangel behoben", "name": NAMEN.get(art, art), "person": "", "alt": zid,
                                   "neu": "", "von": tag.isoformat(), "quelle": "FM-Liste"})
    if n:
        speicher.schreiben(DATEI, {"maengel": [m.zu_dict() for m in alle]})
        protokoll.eintragen(protokolle)
    return n


def _sort(zid: str) -> tuple:
    haus, _, nr = zid.removeprefix("GS-").partition("-")
    try:
        h = float(haus)
    except ValueError:
        h = 99.0
    ziffern = "".join(c for c in nr if c.isdigit())
    return h, nr[:1].isalpha(), int(ziffern or 0), nr
