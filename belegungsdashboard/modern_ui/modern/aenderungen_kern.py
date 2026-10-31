"""Änderungen im System für einen Zeitraum: Zimmerzuteilungen, Umzüge, Auszüge, Mängel – als Tabelle und Excel.

Quelle ist das Änderungsprotokoll des Hauptprogramms (``protokoll.json``), in das beide Oberflächen schreiben: wer
wann welche Person in welches Zimmer geplant, umgeplant oder herausgenommen hat, Umzüge, Abbrüche, gemeldete und
behobene Mängel. Gebucht wird weiterhin im Belegungssystem – diese Liste zeigt, was hier geplant wurde und dort
nachzutragen bzw. zu prüfen ist. Rein, ohne NiceGUI.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from belegung import protokoll

from .standort_kern import kurz

ARTEN = {"alle": "Alle", "zuteilung": "Zuteilungen", "umzug": "Umzüge & Auszüge", "maengel": "Mängel"}


def art_von(aktion: str) -> str:
    a = (aktion or "").lower()
    if a.startswith(("umzug", "abbruch", "vorzeitiger auszug", "auszug")) or a == "sonstiges":
        return "umzug"
    if "mangel" in a or "bewohnbar" in a:
        return "maengel"
    if a in ("geplant", "umgeplant", "entfernt", "rückgängig", "schlüsselliste", "wieder auf geplant") or \
            a.startswith(("planung", "gilt:")):
        return "zuteilung"
    return "sonstiges"


@dataclass
class Aenderung:
    zeit: datetime
    wer: str
    aktion: str
    art: str
    name: str             # Person bzw. bei Mängeln die Mangelart
    alt: str              # Zimmer vorher (Kurzform)
    neu: str              # Zimmer danach (Kurzform)
    ab: date | None       # betroffener Tag (Anreise, Umzugstag, Auszug …)
    quelle: str

    @property
    def zimmer(self) -> str:
        if self.alt and self.neu and self.alt != self.neu:
            return f"{self.alt} → {self.neu}"
        return self.neu or self.alt


def _datum(text) -> date | None:
    try:
        return date.fromisoformat(str(text)[:10])
    except (TypeError, ValueError):
        return None


def liste(von: date, bis: date, art: str = "alle") -> list[Aenderung]:
    """Alle Einträge mit Zeitpunkt zwischen ``von`` und ``bis`` (je einschließlich), neueste zuerst."""
    erg = []
    for e in protokoll.laden():
        try:
            zeit = datetime.fromisoformat(e.get("zeit", ""))
        except (TypeError, ValueError):
            continue
        if not von <= zeit.date() <= bis:
            continue
        a = art_von(e.get("aktion", ""))
        if art != "alle" and a != art:
            continue
        erg.append(Aenderung(zeit, e.get("wer", ""), e.get("aktion", ""), a, e.get("name", ""), kurz(str(e.get("alt") or "")),
                             kurz(str(e.get("neu") or "")), _datum(e.get("von")), e.get("quelle", "")))
    return sorted(erg, key=lambda x: x.zeit, reverse=True)


def zeitraum(vorgabe: str, heute: date | None = None) -> tuple[date, date]:
    heute = heute or date.today()
    montag = heute - timedelta(days=heute.weekday())
    return {"woche": (montag, heute), "vorwoche": (montag - timedelta(days=7), montag - timedelta(days=1)),
            "7": (heute - timedelta(days=6), heute), "30": (heute - timedelta(days=29), heute)}.get(vorgabe, (heute, heute))


def excel(pfad: Path, eintraege: list[Aenderung], von: date, bis: date) -> int:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Änderungen"
    ws.append([f"Änderungen {von:%d.%m.%Y} – {bis:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    kopf = ["Zeitpunkt", "Art", "Aktion", "Person / Mangel", "Zimmer vorher", "Zimmer danach", "Ab", "Quelle", "Wer"]
    ws.append(kopf)
    for c in ws[3]:
        c.font = Font(bold=True)
    for x in eintraege:
        ws.append([x.zeit.replace(microsecond=0), ARTEN.get(x.art, "Sonstiges"), x.aktion, x.name, x.alt, x.neu, x.ab,
                   x.quelle, x.wer])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(ws.max_row, 7).number_format = "DD.MM.YYYY"
    for i, breite in enumerate([17, 16, 26, 30, 14, 14, 12, 24, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(kopf))}{max(ws.max_row, 3)}"
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
    return len(eintraege)


def excel_speichern(von: date, bis: date, art: str = "alle") -> tuple[str, Path]:
    from belegung import speicher

    eintraege = liste(von, bis, art)
    ziel = Path(speicher.einstellungen().get("export_ordner") or speicher.pfad("Exporte"))
    pfad = ziel / f"Aenderungen_{von:%Y-%m-%d}_bis_{bis:%Y-%m-%d}.xlsx"
    n = excel(pfad, eintraege, von, bis)
    return f"{n} Änderungen gespeichert", pfad
