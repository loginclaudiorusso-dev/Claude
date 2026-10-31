"""0.37: Änderungen je Zeitraum (Excel), Schloss/Schlüssel → FM-Auftrag statt Umplanen. Erfundene Daten."""

from datetime import date, datetime, timedelta

import pytest


@pytest.fixture
def leer(tmp_path, monkeypatch):
    from belegung import aenderungen

    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    aenderungen.verlauf_leeren()
    yield tmp_path
    aenderungen.verlauf_leeren()


def test_aenderungen_zeitraum_und_excel(leer):
    import openpyxl

    from belegung import protokoll, speicher
    from modern import aenderungen_kern as ak

    h = date.today()
    alt = (datetime.now() - timedelta(days=20)).isoformat(timespec="seconds")
    speicher.schreiben(protokoll.DATEI, {"eintraege": [
        {"zeit": alt, "wer": "test", "aktion": "geplant", "name": "Alt, Anton", "alt": "", "neu": "GS-2-101", "von": "2030-01-01"}]})
    protokoll.eintragen([
        {"aktion": "geplant", "name": "Neu, Nina", "person": "id:1", "alt": "", "neu": "GS-3.2-111", "von": h.isoformat(), "quelle": "t"},
        {"aktion": "Umzug", "name": "Neu, Nina", "person": "id:1", "alt": "GS-3.2-111", "neu": "GS-3.3-101", "von": h.isoformat(), "quelle": "t"},
        {"aktion": "Mangel gemeldet", "name": "Heizung defekt", "person": "", "alt": "", "neu": "GS-2-104", "von": h.isoformat(), "quelle": "t"}])
    von, bis = ak.zeitraum("7", h)
    alle = ak.liste(von, bis)
    assert [x.art for x in alle] and {x.art for x in alle} == {"zuteilung", "umzug", "maengel"}
    assert all(x.name != "Alt, Anton" for x in alle)                    # außerhalb des Zeitraums
    zut = ak.liste(von, bis, "zuteilung")
    assert [(x.name, x.zimmer) for x in zut] == [("Neu, Nina", "3.2-111")]
    assert ak.liste(von, bis, "umzug")[0].zimmer == "3.2-111 → 3.3-101"
    text, pfad = ak.excel_speichern(von, bis)
    ws = openpyxl.load_workbook(pfad).active
    assert ws["A3"].value == "Zeitpunkt" and ws.max_row == 6 and text.startswith("3 Änderungen")
    m = ak.zeitraum("vorwoche", date(2030, 3, 13))                       # Mittwoch → Mo–So der Vorwoche
    assert m == (date(2030, 3, 4), date(2030, 3, 10))


def _anreise(tage, name="Schloss, Sam", tn="977"):
    from belegung import anreiseliste
    from modern import kern

    an = date.today() + timedelta(days=tage)
    anreiseliste.speichern([anreiseliste.Person(name, "GS RVL", an, True, tn_id=tn, abreise=an + timedelta(days=60))], {})
    kern.vorschlaege_uebernehmen(an.isoformat(), "alle")
    return an


def test_schloss_defekt_fm_statt_umplanen(leer):
    from belegung import maengel
    from belegung import zimmerplan as zp
    from modern import kern

    an = _anreise(10)
    st = zp.laden()
    zid = st.zugewiesen(st.bedarf[0].schluessel)["zimmer"]
    maengel.setzen(zid, "tuer", True, date.today())
    maengel.sperren_setzen(zid, "tuer", True)                           # Schloss defekt → unbewohnbar
    st = zp.laden()
    assert zid in st.lage.unbewohnbar
    assert not kern.problemfaelle(st, date.today())                      # noch 10 Tage: nicht umplanen
    fs = kern.fm_faelle(st, date.today())
    assert [(f.zimmer, f.frist) for f in fs] == [(zid, an)] and not fs[0].knapp
    z = kern.zeilen(st, an.isoformat(), date.today())[0]
    assert z.gespeichert == zid and z.vorschlag is None and "FM-Auftrag" in z.grund
    karten = {k.id: k for k in kern.karten(st, date.today())}
    assert karten["fm"].stufe == "bald" and "unbewohnbar" not in karten
    assert "Zimmer " + kern.kurz(zid) in kern.fm_text(fs) and "Sam" not in kern.fm_text(fs)   # Auftrag ohne Namen
    # drei Tage vorher wird es knapp: dann zurück in „Prüfen und zuweisen“
    assert kern.problemfaelle(st, an - timedelta(days=2))
    assert kern.fm_behoben([zid]) == "1 Zimmer wieder nutzbar"
    assert zid not in zp.laden().lage.unbewohnbar


def test_schimmel_bleibt_umplanen(leer):
    from belegung import maengel
    from belegung import zimmerplan as zp
    from modern import kern

    _anreise(10, "Schimmel, Susi", "978")
    st = zp.laden()
    zid = st.zugewiesen(st.bedarf[0].schluessel)["zimmer"]
    maengel.setzen(zid, "schimmel", True, date.today())
    st = zp.laden()
    assert kern.problemfaelle(st, date.today()) and not kern.fm_faelle(st, date.today())


def test_postfach_auswahl_rueckt_nach():
    """Erledigte Aufgabe verschwindet → die nächste in der alten Reihenfolge wird gewählt (wie im Posteingang)."""
    from types import SimpleNamespace

    from modern import seite_heute as sh

    k = lambda i: SimpleNamespace(id=i)
    a = SimpleNamespace(karten=[k("a"), k("b"), k("c")], pf_auswahl="b", pf_reihe=["a", "b", "c"])
    assert sh.auswahl(a) == "b"
    a.karten = [k("a"), k("c")]
    assert sh.auswahl(a) == "c"
    a.pf_auswahl, a.karten = "c", [k("a")]
    assert sh.auswahl(a) == "a"
    a.karten = []
    assert sh.auswahl(a) == sh.UEBERBLICK
