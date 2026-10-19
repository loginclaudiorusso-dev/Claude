"""6.24: Zimmer von Hand zur Reinigung hinzufügen – Liste, Druck/Excel, Assistenz, Rückgängig."""

from datetime import date

import pytest

from belegung.gebaeudeplan import Belegung

MO = date(2026, 10, 19)


@pytest.fixture
def lage(tmp_path, monkeypatch):
    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    from belegung import zimmerplan as zp

    zimmer = [zp.Zimmer("GS-2-201", "2", "2", "201", 1), zp.Zimmer("GS-2-202", "2", "2", "202", 1)]
    bel = {"GS-2-201": [Belegung("GS-2-201", date(2026, 9, 1), date(2026, 10, 21), "belegt", "Beispiel, Kai"),
                        Belegung("GS-2-201", date(2026, 10, 26), date(2027, 1, 31), "geplant", "Probe, Lena")],
           "GS-2-202": [Belegung("GS-2-202", date(2026, 9, 1), date(2027, 3, 31), "belegt", "Testfall, Hanna")]}
    return zp.Lage({z.id: z for z in zimmer}, bel)


def test_von_hand_auf_der_liste(lage):
    from belegung import reinigungsliste as rl
    from belegung import zimmerplan as zp

    assert [r.zimmer.id for r in zp.reinigungen(lage, MO, MO.replace(day=25))] == ["GS-2-201"]   # nur Auszug
    e = zp.reinigung_hinzufuegen("GS-2-202", date(2026, 10, 22), "Grundreinigung")
    zp.reinigung_hinzufuegen("GS-2-201", date(2026, 10, 21), "")            # gleicher Tag wie der Auszug
    rs = {r.zimmer.id: r for r in zp.reinigungen(lage, MO, MO.replace(day=25))}
    assert set(rs) == {"GS-2-201", "GS-2-202"} and len(zp.reinigungen(lage, MO, MO.replace(day=25))) == 2
    hand = rs["GS-2-202"]
    assert (hand.art, hand.manuell, hand.tag, hand.bleibt) == ("Grundreinigung", "Grundreinigung", date(2026, 10, 22), 1)
    assert rl.auszug_datum(hand) == "Do 22.10. · von Hand" and "Grundreinigung" in rl.details(hand)
    aus = rs["GS-2-201"]                                                     # bleibt eine Zeile: Auszug + Grund
    assert aus.art == "Auszug" and aus.manuell == "Zwischenreinigung" and aus.naechste_anreise == date(2026, 10, 26)
    assert "2 von Hand hinzugefügt" in rl.zaehlung(list(rs.values()))
    html = rl.html_text(list(rs.values()), MO, MO.replace(day=25))
    assert "Grundreinigung" in html and "Zwischenreinigung" in html
    # andere Woche: nicht dabei; entfernen
    assert not any(r.manuell for r in zp.reinigungen(lage, date(2026, 10, 26), date(2026, 11, 1)))
    assert zp.reinigung_entfernen(e["id"]) and "GS-2-202" not in {r.zimmer.id for r in zp.reinigungen(lage, MO, MO.replace(day=25))}


def test_excel_mit_notiz(lage, tmp_path):
    from openpyxl import load_workbook

    from belegung import reinigungsliste as rl
    from belegung import zimmerplan as zp

    zp.reinigung_hinzufuegen("GS-2-202", date(2026, 10, 20), "Nach Reparatur / Handwerker")
    pfad = rl.excel_speichern(tmp_path / "r.xlsx", zp.reinigungen(lage, MO, MO.replace(day=25)), MO, MO.replace(day=25))
    werte = [c.value for z in load_workbook(pfad).active.iter_rows(min_row=4) for c in z]
    assert "Nach Reparatur / Handwerker" in werte


def test_assistenz_setzt_auf_die_reinigungsliste(lage):
    from belegung import assistenz
    from belegung import zimmerplan as zp

    v = assistenz.verstehen("2-202 Zwischenreinigung morgen", {"GS-2-201", "GS-2-202"}, MO)
    arten = {a.art: a.an for a in v.aktionen}
    assert arten["reinigung"] and not arten.get("mangel", False)
    e = assistenz.ausfuehren([v], MO)
    assert [(r["zimmer"], r["tag"], r["grund"]) for r in zp.manuelle_reinigungen()] == \
        [("GS-2-202", "2026-10-20", "Zwischenreinigung")]
    assistenz.rueckgaengig(e["id"])
    assert zp.manuelle_reinigungen() == []


def test_zimmer_eingabe_im_dialog():
    from belegung.ui.reinigung_dialog import zimmer_lesen

    bekannt = {"GS-2-204", "GS-2-205", "GS-6-403", "GS-2-403", "GS-3.1-E02"}
    assert zimmer_lesen("2-204, 2-205", bekannt) == (["GS-2-204", "GS-2-205"], [])
    assert zimmer_lesen("204 E02", bekannt) == (["GS-2-204", "GS-3.1-E02"], [])
    assert zimmer_lesen("Haus 6 Zimmer 403", bekannt) == (["GS-6-403"], [])
    ids, hinweise = zimmer_lesen("403", bekannt)
    assert ids == [] and "mehreren Häusern" in hinweise[0]
    assert "gibt es im Zimmerplan nicht" in zimmer_lesen("999", bekannt)[1][0]


def test_knopf_im_reiter_oeffnet_die_maske(tmp_path, monkeypatch, qapp_klick=None):
    """Regression 6.24.0: Ein echter Klick übergab False als Zimmer – die Maske ging nicht auf."""
    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication, QPushButton

    app = QApplication.instance() or QApplication([])
    from belegung.ui import reinigung_dialog as rd
    from belegung.ui.basis import Zustand
    from belegung.ui.seiten.zimmerplan import ZimmerplanSeite

    geoeffnet = []
    monkeypatch.setattr(rd.ReinigungDialog, "exec", lambda self: geoeffnet.append(self.feld.text()) or 0)
    seite = ZimmerplanSeite(Zustand())
    seite.show()
    seite.aktualisieren()
    app.processEvents()
    knopf = next(b for b in seite.findChildren(QPushButton) if b.text().strip() == "Zimmer hinzufügen …")
    QTest.mouseClick(knopf, Qt.LeftButton)
    assert geoeffnet == [""]
    seite.close()
