"""6.19: „!“-Filter und Mängelbericht für das FM (Druck, Excel zum Abarbeiten, Erledigte wieder einlesen)."""

from datetime import date

import pytest

HEUTE = date(2026, 10, 17)


@pytest.fixture
def daten(tmp_path, monkeypatch):
    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    from belegung import maengel as mg

    mg.setzen("GS-3.3-101", "moebel", True, date(2026, 4, 16), "Nachttisch Loch")
    mg.setzen("GS-2-402", "tuer", True, date(2026, 3, 6), "Schloss schwergängig")
    mg.setzen("GS-2-402", "bad", True, date(2026, 3, 6), "WC-Deckel lose")
    return tmp_path


def test_bericht_sortiert_ohne_namen(daten):
    from belegung import maengel as mg

    z = mg.bericht(mg.offen(), status={"GS-2-402": "belegt"}, heute=HEUTE)
    assert [(r.kurz, r.mangel) for r in z] == [("2-402", "Bad/Dusche/WC defekt"), ("2-402", "Tür/Schloss/Schlüssel defekt"),
                                              ("3.3-101", "Möbel beschädigt")]
    assert z[0].zimmer_heute == "belegt" and z[0].tage == 225 and z[0].details == "WC-Deckel lose"
    html = mg.fm_html(mg.offen(), status={"GS-2-402": "belegt"}, heute=HEUTE)
    assert "3 offene Mängel in 2 Zimmern" in html and "Haus 2" in html and "Haus 3.3" in html


def test_excel_abarbeiten_und_wieder_einlesen(daten, tmp_path):
    from openpyxl import load_workbook

    from belegung import maengel as mg

    pfad = tmp_path / "fm.xlsx"
    assert mg.fm_excel(pfad, mg.offen(), heute=HEUTE) == 3
    wb = load_workbook(pfad)
    ws = wb["Mängel"]
    assert [c.value for c in ws[1]] == mg.EXCEL_SPALTEN
    ws["I2"] = date(2026, 10, 20)          # das FM trägt ein: Bad erledigt
    ws["J2"] = "Deckel getauscht"
    ws["I3"] = "x"                          # Tür: „x“ = heute erledigt
    ws["I4"] = "irgendwann"                 # unlesbar → von Hand
    wb.save(pfad)
    eintraege, fehler = mg.fm_erledigte_lesen(pfad, heute=HEUTE)
    assert eintraege == [("GS-2-402", "bad", date(2026, 10, 20), "Deckel getauscht"), ("GS-2-402", "tuer", HEUTE, "")]
    assert fehler == ["3.3-101: „irgendwann“"]
    assert mg.erledigte_uebernehmen(eintraege) == 2
    assert [m.art for m in mg.offen()] == ["moebel"]
    bad = next(m for m in mg.laden() if m.art == "bad")
    assert bad.behoben == date(2026, 10, 20) and bad.notiz == "WC-Deckel lose · FM: Deckel getauscht"
    assert mg.erledigte_uebernehmen(eintraege) == 0          # zweimal einlesen ändert nichts


def test_fremde_excel_wird_erkannt(tmp_path):
    from openpyxl import Workbook

    from belegung import maengel as mg

    wb = Workbook()
    wb.active.append(["irgendwas", "anderes"])
    wb.save(tmp_path / "x.xlsx")
    assert mg.fm_erledigte_lesen(tmp_path / "x.xlsx")[0] == []


def test_nur_unbewohnbare_maengel_sperren_die_zuteilung(tmp_path, monkeypatch):
    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    from belegung import maengel as mg
    from belegung import zimmerplan as zp

    st = zp.laden()
    z = next(x for x in st.zimmer if x.aktiv and zp.erlaubt(x, "Reha") and x.betten == 1 and not x.nur_maenner
             and not x.tiere and not zp.nur_kurz(x))
    b = zp.Bedarf("p:Testfall", "Testfall, Hanna", "GS ASS XX 0000", "Reha", date(2027, 3, 1), date(2027, 6, 30), "w")
    mg.setzen(z.id, "fenster", True, notiz="Plissee defekt")          # kein Grund
    assert zp.passt(zp.laden().lage, z, b) is None
    mg.setzen(z.id, "matratze", True, notiz="Matratze fehlt")
    assert mg.sperren_setzen(z.id, "matratze", True)                  # Nutzer: macht unbewohnbar
    lage = zp.laden().lage
    assert lage.unbewohnbar == {z.id: "Matratze erneuern (Matratze fehlt)"}
    assert zp.passt(lage, z, b).startswith(zp.UNBEWOHNBAR_GRUND)
    vorschlag = zp.vorschlagen(lage, [b])[0]
    assert vorschlag.zimmer is not None and vorschlag.zimmer.id != z.id
    fest = zp.vorschlagen(lage, [b], {b.schluessel: z.id})[0]          # von Hand: geht, mit Warnung
    assert fest.zimmer.id == z.id and any("nicht bewohnbar" in w for w in fest.warnungen)
    mg.setzen(z.id, "matratze", False)                                 # behoben → wieder frei
    assert zp.passt(zp.laden().lage, z, b) is None
    mg.setzen(z.id, "schimmel", True)                                  # Schimmel sperrt von sich aus
    assert z.id in zp.laden().lage.unbewohnbar
    assert any(r.dringend for r in mg.bericht(mg.offen()))


def test_stufen_und_filter(tmp_path, monkeypatch):
    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    from belegung import maengel as mg
    from belegung import zimmerliste as zl
    from belegung import zimmerplan as zp

    mg.setzen("GS-2-101", "fenster", True)              # Rollo/Plissee: leicht (!)
    mg.setzen("GS-2-102", "fenster", True)
    mg.setzen("GS-2-102", "heizung", True)              # dringend (!!) – höchste Stufe zählt
    mg.setzen("GS-2-103", "moebel", True)
    mg.sperren_setzen("GS-2-103", "moebel", True)       # vom Nutzer als unbewohnbar markiert (!!!)
    je = mg.offen_je_zimmer()
    assert [mg.stufe(je[z]) for z in ("GS-2-101", "GS-2-102", "GS-2-103")] == [1, 2, 3]
    assert [mg.STUFEN[s][3] for s in (1, 2, 3)] == ["gelb", "orange", "rot"] and mg.stufe([]) == 0
    assert mg.stufe_text(3) == "rot – nicht bewohnbar" and mg.zeichen(2) == "!"
    st = zp.laden()
    ids = lambda f: {r.zimmer.id for r in zl.filtern(st, f)}
    assert ids(zl.Filter(maengel={"stufe:3"})) == {"GS-2-103"}
    assert ids(zl.Filter(maengel={"stufe:2", "stufe:3"})) == {"GS-2-102", "GS-2-103"}
    assert ids(zl.Filter(maengel={"*"})) == {"GS-2-101", "GS-2-102", "GS-2-103"}
    assert "rot (nicht bewohnbar)" in zl.Filter(maengel={"stufe:3"}).beschreibung()
    z = mg.bericht(mg.offen())
    assert [(r.kurz, r.stufe) for r in z][:3] == [("2-101", 1), ("2-102", 2), ("2-102", 1)]   # je Zimmer: wichtigster zuerst
