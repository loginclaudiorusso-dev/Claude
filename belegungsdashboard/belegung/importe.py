"""Import von Excel-/CSV-Listen für Mieten und Anreisen.

Die Listen kommen aus unterschiedlichen Quellen und sind nicht einheitlich aufgebaut.
Deshalb: Kopfzeile automatisch suchen, Spalten per Synonym erraten, Nutzer bestätigt bzw.
korrigiert die Zuordnung in der Vorschau. Die Zuordnung wird je Dateiname gemerkt, sodass
ein erneuter Upload derselben Liste ohne Rückfragen durchläuft.

Ein Import ersetzt den vorherigen Datensatz derselben Art vollständig (Liste = aktueller
Stand), damit sich nichts doppelt aufsummiert.
"""

from __future__ import annotations

import csv
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from . import speicher
from .konfig import STANDARD_STANDORT_MIETER, STANDORTE, standort_normalisieren
from .pivot import EXCEL_EPOCH

ARTEN = {"mieten": "Mieten", "anreisen": "Anreisen"}
# Über den allgemeinen Listen-Import kommen nur noch Mieten; Anreisen (Goslar) und UWT
# haben eigene Leser (anreiseliste.py, uwt.py).
LISTEN_ARTEN = ["mieten"]
FELDER = ["von", "bis", "standort", "anzahl", "bezeichnung", "kategorie"]
FELD_LABEL = {
    "von": "Von / Anreise", "bis": "Bis / Abreise", "standort": "Standort", "anzahl": "Anzahl",
    "bezeichnung": "Bezeichnung / Name", "kategorie": "Kategorie",
}
OFFENES_ENDE = date(2099, 12, 31)

_SYNONYME: dict[str, list[str]] = {
    "von": ["von", "beginn", "start", "ab", "anreise", "anreisedatum", "einzug", "mietbeginn",
            "vertragsbeginn", "eintritt", "aufnahme", "ankunft", "massnahmebeginn", "datum"],
    "bis": ["bis", "ende", "abreise", "abreisedatum", "auszug", "mietende", "vertragsende",
            "austritt", "entlassung", "enddatum", "massnahmeende"],
    "standort": ["standort", "ort", "bfw", "einrichtung", "niederlassung", "liegenschaft", "betrieb"],
    "anzahl": ["anzahl", "personen", "plaetze", "platze", "betten", "menge", "tn", "pax", "anzahlpersonen"],
    "bezeichnung": ["mieter", "name", "nachname", "firma", "kunde", "gast", "teilnehmer", "bezeichnung",
                    "vertrag", "beschreibung", "bemerkung"],
    "kategorie": ["kategorie", "typ", "art", "vertragsart", "kostentraeger", "kostentrager"],
}


def _norm(text: object) -> str:
    t = str(text or "").lower()
    t = t.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    return re.sub(r"[^a-z]", "", t)


@dataclass
class Zuordnung:
    sheet: str | None = None
    kopfzeile: int = 0                      # 0-basierter Zeilenindex
    spalten: dict[str, int | None] = field(default_factory=lambda: {f: None for f in FELDER})
    standard_standort: str | None = None    # für Zeilen ohne erkennbaren Standort
    dauer_monate: int = 0                   # Anreisen ohne Abreise: angenommene Dauer (0 = nur Termin)

    def to_dict(self) -> dict:
        return asdict(self)

    @staticmethod
    def from_dict(d: dict) -> "Zuordnung":
        z = Zuordnung()
        for k, v in d.items():
            if hasattr(z, k):
                setattr(z, k, v)
        z.spalten = {f: z.spalten.get(f) for f in FELDER}
        return z


@dataclass
class Tabelle:
    pfad: Path
    sheets: list[str]
    sheet: str | None
    zeilen: list[list]


def tabelle_lesen(pfad: Path, sheet: str | None = None) -> Tabelle:
    endung = pfad.suffix.lower()
    if endung in (".csv", ".txt"):
        roh = pfad.read_bytes()
        text = None
        for kodierung in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = roh.decode(kodierung)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialekt = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        except csv.Error:
            class dialekt(csv.excel):
                delimiter = ";"
        zeilen = [list(z) for z in csv.reader(text.splitlines(), dialekt)]
        return Tabelle(pfad, [], None, zeilen)
    if endung == ".xls":
        raise ValueError("Das alte .xls-Format wird nicht unterstützt – bitte in Excel als .xlsx speichern.")

    import openpyxl
    wb = openpyxl.load_workbook(pfad, data_only=True, read_only=True)
    try:
        namen = list(wb.sheetnames)
        ws = wb[sheet] if sheet in namen else wb[_bestes_sheet(wb, namen)]
        zeilen = [list(z) for z in ws.iter_rows(values_only=True)]
        return Tabelle(pfad, namen, ws.title, _leere_raender_entfernen(zeilen))
    finally:
        wb.close()


def _bestes_sheet(wb, namen: list[str]) -> str:
    """Sheet mit der am besten erkennbaren Kopfzeile, sonst das erste."""
    bestes, beste_punkte = namen[0], -1
    for name in namen:
        zeilen = [list(z) for _, z in zip(range(30), wb[name].iter_rows(values_only=True))]
        punkte = max((_kopf_punkte(z) for z in zeilen), default=0)
        if punkte > beste_punkte:
            bestes, beste_punkte = name, punkte
    return bestes


def _leere_raender_entfernen(zeilen: list[list]) -> list[list]:
    zeilen = [z for z in zeilen]
    while zeilen and all(v in (None, "") for v in zeilen[-1]):
        zeilen.pop()
    breite = max((max((i + 1 for i, v in enumerate(z) if v not in (None, "")), default=0) for z in zeilen), default=0)
    return [list(z[:breite]) + [None] * (breite - len(z)) for z in zeilen]


def _feld_fuer_kopf(kopf: object) -> tuple[str | None, float]:
    n = _norm(kopf)
    if not n:
        return None, 0
    bestes, punkte = None, 0.0
    for feld, synonyme in _SYNONYME.items():
        for rang, s in enumerate(synonyme):
            if n == s:
                p = 2.0
            elif n.startswith(s) or (len(s) >= 4 and s in n):
                p = 1.5
            else:
                continue
            p += 1.0 - rang / len(synonyme)  # frühere Synonyme sind typischer
            if p > punkte:
                bestes, punkte = feld, p
    return bestes, punkte


def _kopf_punkte(zeile: list) -> int:
    return sum(1 for v in zeile if isinstance(v, str) and _feld_fuer_kopf(v)[0])


def kopfzeile_finden(zeilen: list[list]) -> int:
    beste, beste_punkte = 0, 0
    for i, z in enumerate(zeilen[:40]):
        p = _kopf_punkte(z)
        if p > beste_punkte:
            beste, beste_punkte = i, p
    return beste


def zuordnung_erraten(tabelle: Tabelle) -> Zuordnung:
    kopf_idx = kopfzeile_finden(tabelle.zeilen)
    kopf = tabelle.zeilen[kopf_idx] if tabelle.zeilen else []
    kandidaten: list[tuple[float, str, int]] = []
    for i, v in enumerate(kopf):
        feld, punkte = _feld_fuer_kopf(v)
        if feld:
            kandidaten.append((punkte, feld, i))
    spalten: dict[str, int | None] = {f: None for f in FELDER}
    for _, feld, i in sorted(kandidaten, key=lambda k: (-k[0], k[2])):
        if spalten[feld] is None and i not in spalten.values():
            spalten[feld] = i
    return Zuordnung(sheet=tabelle.sheet, kopfzeile=kopf_idx, spalten=spalten)


def datum_lesen(wert) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, (int, float)) and not isinstance(wert, bool) and 20000 < wert < 80000:
        return (EXCEL_EPOCH + timedelta(days=float(wert))).date()
    if isinstance(wert, str):
        t = wert.strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(t, fmt).date()
            except ValueError:
                continue
    return None


def _monate_addieren(d: date, monate: int) -> date:
    import calendar
    m = d.month - 1 + monate
    jahr, monat = d.year + m // 12, m % 12 + 1
    return date(jahr, monat, min(d.day, calendar.monthrange(jahr, monat)[1]))


def _kategorie(art: str, wert) -> str:
    if art == "anreisen":
        return "Anreise"
    n = _norm(wert)
    if n.startswith("gast") or n.startswith("gaest"):
        return "Gäste"
    return "Mieter"


def umwandeln(tabelle: Tabelle, zuordnung: Zuordnung, art: str) -> tuple[list[dict], list[str]]:
    """Datenzeilen -> Einträge {kategorie, standort, von, bis, anzahl, bezeichnung, nur_termin}.
    Gibt zusätzlich eine Liste verständlicher Hinweise zu übersprungenen Zeilen zurück."""
    sp = zuordnung.spalten
    if sp.get("von") is None:
        raise ValueError("Bitte die Spalte für 'Von / Anreise' zuordnen.")
    standard_standort = zuordnung.standard_standort or (STANDARD_STANDORT_MIETER if art == "mieten" else None)

    def zelle(z: list, feld: str):
        i = sp.get(feld)
        return z[i] if i is not None and i < len(z) else None

    eintraege, hinweise = [], []
    for nr, z in enumerate(tabelle.zeilen[zuordnung.kopfzeile + 1:], start=zuordnung.kopfzeile + 2):
        if all(v in (None, "") for v in z):
            continue
        von = datum_lesen(zelle(z, "von"))
        if von is None:
            hinweise.append(f"Zeile {nr}: kein gültiges Datum in 'Von' – übersprungen.")
            continue
        bis_roh = zelle(z, "bis")
        bis = datum_lesen(bis_roh)
        nur_termin = False
        if bis is None:
            if art == "mieten":
                bis = OFFENES_ENDE
            elif zuordnung.dauer_monate > 0:
                bis = _monate_addieren(von, zuordnung.dauer_monate) - timedelta(days=1)
            else:
                bis, nur_termin = von, True
        if bis < von:
            hinweise.append(f"Zeile {nr}: 'Bis' liegt vor 'Von' – übersprungen.")
            continue
        standort_roh = zelle(z, "standort")
        standort = standort_normalisieren(standort_roh) or standard_standort
        if standort is None:
            hinweise.append(f"Zeile {nr}: Standort '{standort_roh or ''}' nicht erkannt – übersprungen.")
            continue
        anzahl_roh = zelle(z, "anzahl")
        try:
            anzahl = int(float(str(anzahl_roh).replace(",", "."))) if anzahl_roh not in (None, "") else 1
        except ValueError:
            hinweise.append(f"Zeile {nr}: Anzahl '{anzahl_roh}' ist keine Zahl – als 1 gewertet.")
            anzahl = 1
        if anzahl <= 0:
            continue
        eintraege.append({
            "kategorie": _kategorie(art, zelle(z, "kategorie")),
            "standort": standort,
            "von": von.isoformat(),
            "bis": bis.isoformat(),
            "anzahl": anzahl,
            "bezeichnung": str(zelle(z, "bezeichnung") or "").strip(),
            "nur_termin": nur_termin,
        })
    return eintraege, hinweise


# ---------------------------------------------------------------------------------------
# Ablage
# ---------------------------------------------------------------------------------------

def _datei(art: str) -> str:
    return f"{speicher.IMPORT_ORDNER}/{art}.json"


def datensatz_laden(art: str) -> dict | None:
    return speicher.lesen(_datei(art), None)


def datensatz_speichern(art: str, quelle: Path, zuordnung: Zuordnung, eintraege: list[dict]) -> dict:
    datensatz = {
        "art": art,
        "quelle": quelle.name,
        "importiert_am": datetime.now().isoformat(timespec="seconds"),
        "eintraege": eintraege,
    }
    speicher.schreiben(_datei(art), datensatz)
    gemerkt = speicher.lesen(f"{speicher.IMPORT_ORDNER}/zuordnungen.json", {})
    gemerkt[f"{art}:{quelle.name.lower()}"] = zuordnung.to_dict()
    speicher.schreiben(f"{speicher.IMPORT_ORDNER}/zuordnungen.json", gemerkt)
    return datensatz


def datensatz_loeschen(art: str) -> None:
    speicher.pfad(_datei(art)).unlink(missing_ok=True)


def gemerkte_zuordnung(art: str, dateiname: str) -> Zuordnung | None:
    gemerkt = speicher.lesen(f"{speicher.IMPORT_ORDNER}/zuordnungen.json", {})
    d = gemerkt.get(f"{art}:{dateiname.lower()}")
    return Zuordnung.from_dict(d) if d else None


def standorte_in(eintraege: list[dict]) -> dict[str, int]:
    return {s: sum(e["anzahl"] for e in eintraege if e["standort"] == s) for s in STANDORTE}
