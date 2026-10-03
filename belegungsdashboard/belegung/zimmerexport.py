"""Excel-Export der Zimmerplanung: Zimmerliste je Anreise, Reinigungsplan je Flur, freie Zimmer."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from .zimmerplan import GRUPPE_LABEL, Lage, Zuteilung, etage_text

_AKZENT = "00908F"
_KOPF_TEXT = "FFFFFF"
_ZEBRA = "F2F7F7"
_WARN = "FFF4E5"
_FEHLER = "FDECEC"
_RAND = "D5DEDE"


def _tabelle(ws, titel: str, untertitel: str, kopf: list[str], zeilen: list[list], breiten: list[int],
             markierung: dict[int, str] | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws["A1"] = titel
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = untertitel
    ws["A2"].font = Font(size=9, color="666666")
    rand = Side(style="thin", color=_RAND)
    for i, k in enumerate(kopf, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color=_KOPF_TEXT)
        c.fill = PatternFill("solid", fgColor=_AKZENT)
        c.alignment = Alignment(vertical="center")
        c.border = Border(bottom=rand)
    for r, zeile in enumerate(zeilen, start=5):
        farbe = (markierung or {}).get(r - 5) or (_ZEBRA if (r - 5) % 2 else None)
        for i, wert in enumerate(zeile, start=1):
            c = ws.cell(row=r, column=i, value=wert)
            if isinstance(wert, (date, datetime)):
                c.number_format = "DD.MM.YYYY"
            c.alignment = Alignment(vertical="top", wrap_text=i == len(zeile))
            c.border = Border(bottom=rand)
            if farbe:
                c.fill = PatternFill("solid", fgColor=farbe)
    for i, b in enumerate(breiten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A5"
    if zeilen:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(kopf))}{4 + len(zeilen)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Seite &P von &N"
    ws.oddFooter.right.text = "&D"


def _geschlecht(z: Zuteilung) -> str:
    g = {"m": "m", "w": "w", "d": "d"}.get(z.bedarf.geschlecht, "?")
    return g + ("*" if z.bedarf.geschlecht_geschaetzt and z.bedarf.geschlecht else "")


def exportieren(pfad: Path, titel: str, zuteilungen: list[Zuteilung], lage: Lage, stichtag: date) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Zimmerliste"
    zeilen, markierung = [], {}
    sortiert = sorted(zuteilungen, key=lambda z: (z.bedarf.von, z.zimmer is None,
                                                  (z.zimmer.haus, z.zimmer.etage, z.zimmer.flur, z.zimmer.nummer) if z.zimmer else (),
                                                  z.bedarf.name))
    for i, zt in enumerate(sortiert):
        b, z = zt.bedarf, zt.zimmer
        hinweis = "; ".join(([] if z else [zt.grund]) + zt.warnungen)
        zeilen.append([i + 1, b.name, b.massnahme, GRUPPE_LABEL.get(b.gruppe, b.gruppe), _geschlecht(zt),
                       f"Haus {z.haus}" if z else "–", etage_text(z.etage) if z else "", z.flur if z else "",
                       z.nr if z else "NICHT ZUGETEILT", "Doppel" if z and z.betten > 1 else "", b.von, b.bis,
                       "ja" if b.tier else "", b.bemerkung, hinweis])
        if z is None:
            markierung[i] = _FEHLER
        elif zt.warnungen:
            markierung[i] = _WARN
    zugeteilt = sum(1 for z in zuteilungen if z.zimmer)
    _tabelle(ws, titel, f"Internat Goslar · {zugeteilt} von {len(zuteilungen)} Personen zugeteilt · erstellt "
                        f"{datetime.now():%d.%m.%Y %H:%M} · * Geschlecht aus dem Vornamen geschätzt",
             ["Nr", "Name", "Maßnahme", "Gruppe", "m/w", "Haus", "Etage", "Flur", "Zimmer", "Typ", "Anreise", "Abreise",
              "Tier", "Bemerkung", "Hinweis"], zeilen, [5, 28, 18, 13, 6, 9, 8, 8, 14, 8, 11, 11, 6, 30, 40], markierung)

    # Reinigung / Vorbereitung je Flur
    ws2 = wb.create_sheet("Je Flur")
    flure: dict[tuple, list[Zuteilung]] = {}
    for zt in sortiert:
        if zt.zimmer:
            flure.setdefault((zt.zimmer.haus, zt.zimmer.etage, zt.zimmer.flur), []).append(zt)
    zeilen2 = []
    for (haus, etage, flur), zts in flure.items():
        zimmer = sorted({zt.zimmer.nr for zt in zts}, key=lambda n: (len(n), n))
        anreise = min(zt.bedarf.von for zt in zts)
        abreise = max(zt.bedarf.bis for zt in zts)
        zeilen2.append([f"Haus {haus}", etage_text(etage), flur, len(zts), len(zimmer), anreise, abreise, ", ".join(zimmer)])
    _tabelle(ws2, f"{titel} – je Flur", "Für Vorbereitung und Reinigung zum Start und Ende",
             ["Haus", "Etage", "Flur", "Personen", "Zimmer", "Anreise", "Abreise bis", "Zimmernummern"],
             zeilen2, [9, 8, 8, 10, 9, 11, 12, 60])

    # Freie Zimmer am Stichtag
    ws3 = wb.create_sheet("Freie Zimmer")
    zeilen3 = []
    for z in sorted(lage.zimmer.values(), key=lambda z: (float(z.haus), z.etage.zfill(2) if z.etage.isdigit() else z.etage, z.flur, z.nummer)):
        if not z.aktiv:
            continue
        status, _ = lage.status_am(z.id, stichtag)
        if status not in ("frei", "teilweise"):
            continue
        frei_bis = lage.frei_bis(z.id, stichtag)
        merkmale = ", ".join(t for t, an in (("Doppel", z.betten > 1), ("Tiere", z.tiere), ("nur Männer", z.nur_maenner),
                                             ("Gäste", z.gaeste)) if an)
        zeilen3.append([f"Haus {z.haus}", etage_text(z.etage), z.flur, z.nr,
                        lage.freie_betten(z.id, stichtag, stichtag), frei_bis or "offen", merkmale, z.notiz])
    _tabelle(ws3, f"Freie Zimmer am {stichtag:%d.%m.%Y}", "Nach Abzug von Belegung, Sperrungen und geplanten Zuweisungen",
             ["Haus", "Etage", "Flur", "Zimmer", "Betten frei", "frei bis", "Merkmale", "Notiz"],
             zeilen3, [9, 8, 8, 10, 11, 12, 22, 40])

    pfad = Path(pfad)
    wb.save(pfad)
    return pfad
