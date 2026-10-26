"""0.18: Zimmer drucken / als Excel je Filter – mit „frei ab“, „frei bis“ und freien Tagen (erfundene Namen)."""

from datetime import date, timedelta

import pytest


@pytest.fixture
def stand(tmp_path, monkeypatch):
    from belegung import aenderungen
    from belegung import gebaeudeplan as gp
    from belegung import zimmerplan as zp

    monkeypatch.setenv("BELEGUNG_DATEN", str(tmp_path))
    aenderungen.verlauf_leeren()
    heute = date(2030, 3, 2)
    st = zp.laden()
    einzel = [z for z in st.zimmer if z.aktiv and z.betten == 1][:3]
    a, b, c = (z.id for z in einzel)
    st.lage.belegungen.setdefault(a, []).append(gp.Belegung(a, heute - timedelta(days=10), heute + timedelta(days=5), "belegt",
                                                            "Muster, Max"))
    st.lage.belegungen.setdefault(a, []).append(gp.Belegung(a, heute + timedelta(days=20), heute + timedelta(days=40), "geplant",
                                                            "Beispiel, Berta"))
    st.lage.belegungen.setdefault(b, []).append(gp.Belegung(b, heute + timedelta(days=9), heute + timedelta(days=30), "geplant",
                                                            "Probe, Paul"))
    st.lage.belegungen.setdefault(c, []).append(gp.Belegung(c, heute - timedelta(days=1), heute + timedelta(days=3), "gesperrt",
                                                            grund="Renoviert"))
    yield st, heute, einzel
    aenderungen.verlauf_leeren()


def test_frei_ab_und_frei_bis(stand):
    from modern import zimmer_kern as zk

    st, heute, (za, zb, zc) = stand
    zeilen = {z.zimmer: z for z in zk.druck_zeilen(st, heute)}
    a, b, c = zeilen[za.kurz], zeilen[zb.kurz], zeilen[zc.kurz]
    p = st.lage.puffer
    assert a.status == "belegt" and a.bewohner == "Muster, Max" and a.belegt_bis == heute + timedelta(days=5)
    assert a.frei_ab == heute + timedelta(days=5 + p) and a.frei_bis == heute + timedelta(days=20 - p)
    assert a.frei_tage == (a.frei_bis - a.frei_ab).days + 1
    assert b.status == "frei" and b.frei_ab == heute and b.frei_bis == heute + timedelta(days=9 - p)
    assert b.frei_tage == 9 - p + 1
    assert c.status == "gesperrt" and c.frei_ab is None and "Renoviert" in c.bewohner
    offen = [z for z in zeilen.values() if z.status == "frei" and z.frei_bis is None]
    assert offen and all(z.frei_tage is None for z in offen)


def test_kachel_und_filter(stand):
    from modern import zimmer_kern as zk

    st, heute, (za, zb, zc) = stand
    frei = zk.druck_zeilen(st, heute, kachel="frei")
    assert frei and all(z.status in ("frei", "teilweise") for z in frei) and za.kurz not in {z.zimmer for z in frei}
    assert [z.zimmer for z in zk.druck_zeilen(st, heute, kachel="belegt")] == [za.kurz]
    assert [z.zimmer for z in zk.druck_zeilen(st, heute, kachel="gesperrt")] == [zc.kurz]
    f = zk.filter_bauen(za.haus, "belegt")
    assert [z.zimmer for z in zk.druck_zeilen(st, heute, f)] == [za.kurz]
    assert all(z.bewohner == "" for z in zk.druck_zeilen(st, heute, kachel="belegt", namen=False))


def test_druck_html(stand):
    from modern import zimmer_kern as zk

    st, heute, (za, _zb, _zc) = stand
    zeilen = zk.druck_zeilen(st, heute, kachel="frei")
    html = zk.druck_html(zeilen, heute, kachel="frei")
    assert "frei bis" in html and "offen" in html
    assert f"Stand {heute:%d.%m.%Y}" in html and "alle Häuser · frei" in html
    ohne = zk.druck_html(zk.druck_zeilen(st, heute, namen=False), heute, namen=False)
    assert "Muster" not in ohne and "<th>Bewohner</th>" not in ohne
    assert "Muster, Max" in zk.druck_html(zk.druck_zeilen(st, heute), heute)


def test_druck_excel(stand, tmp_path):
    from openpyxl import load_workbook

    from modern import zimmer_kern as zk

    st, heute, (za, zb, _zc) = stand
    zeilen = zk.druck_zeilen(st, heute, zk.filter_bauen(za.haus))
    text, pfad = zk.druck_excel(zeilen, heute, zk.filter_bauen(za.haus))
    assert pfad.exists() and str(tmp_path) in str(pfad) and "Zimmer" in text
    ws = load_workbook(pfad).active
    kopf = [c.value for c in ws[4]]
    assert {"frei ab", "frei bis", "freie Tage", "Bewohner"} <= set(kopf)
    sp = kopf.index("frei bis")
    werte = {r[kopf.index("Zimmer")].value: r[sp].value for r in ws.iter_rows(min_row=5)}
    p = st.lage.puffer
    assert werte[za.kurz].date() == heute + timedelta(days=20 - p)
    assert "offen" in werte.values()
    _t, pfad2 = zk.druck_excel(zk.druck_zeilen(st, heute, kachel="frei", namen=False), heute, kachel="frei", namen=False)
    assert "frei" in pfad2.name and "ohne_Namen" in pfad2.name
    assert "Bewohner" not in [c.value for c in load_workbook(pfad2).active[4]]


def test_kachel_mehrfach_und_mangelart():
    """0.19: mehrere Kacheln – Status mit „oder“, Mängel/unbewohnbar schränken ein; Mangelart gilt zusätzlich."""
    from modern import zimmer_kern as zk

    assert zk.kacheln(None) == frozenset() and zk.kacheln("frei") == {"frei"} and zk.kacheln(["frei", "geplant"]) == {"frei", "geplant"}
    p = zk.kachel_passt
    assert p("frei", 0, False, set(), frozenset())                                  # nichts gewählt = alles
    assert p("geplant", 0, False, set(), {"frei", "geplant"}) and p("teilweise", 0, False, set(), {"frei", "geplant"})
    assert not p("belegt", 0, False, set(), {"frei", "geplant"})
    assert p("frei", 2, False, {"heizung"}, {"frei", "maengel"}) and not p("frei", 0, False, set(), {"frei", "maengel"})
    assert not p("belegt", 2, False, {"heizung"}, {"frei", "maengel"})
    assert p("belegt", 2, False, {"heizung"}, frozenset(), ["heizung"]) and not p("belegt", 2, False, {"wand"}, frozenset(), ["heizung"])
    assert p("frei", 3, True, {"schimmel"}, {"frei"}, ["stufe:3"]) and not p("frei", 2, False, {"heizung"}, {"frei"}, ["stufe:3"])
    assert not p("frei", 2, False, {"heizung"}, {"frei"}, ["stufe:3", "heizung"])       # Stufe und Art: beides muss passen
    assert zk.kachel_text({"geplant", "frei", "maengel"}, ["heizung"]) == "frei oder geplant · mit Mängeln · Mangel: Heizung defekt"


def test_druck_mit_mangelart(stand):
    from belegung import maengel as mg
    from modern import zimmer_kern as zk

    st, heute, (za, zb, zc) = stand
    mg.setzen(zb.id, "heizung", True, heute)
    mg.setzen(za.id, "wand", True, heute)
    assert [z.zimmer for z in zk.druck_zeilen(st, heute, maengel=["heizung"])] == [zb.kurz]
    assert {z.zimmer for z in zk.druck_zeilen(st, heute, kachel={"belegt", "maengel"})} == {za.kurz}
    assert {z.zimmer for z in zk.druck_zeilen(st, heute, kachel={"belegt", "gesperrt"})} == {za.kurz, zc.kurz}
    zeilen = zk.druck_zeilen(st, heute, kachel={"frei"}, maengel=["heizung"])
    html = zk.druck_html(zeilen, heute, kachel={"frei"}, maengel=["heizung"])
    assert "frei · Mangel: Heizung defekt" in html and zb.kurz in html and za.kurz not in html
    _t, pfad = zk.druck_excel(zeilen, heute, kachel={"frei"}, maengel=["heizung"])
    assert pfad.name.startswith("Zimmer_frei_Maengelart_")
    karten = {k.zimmer: k for ks in zk.uebersicht(st, heute)[zb.haus].values() for k in ks}
    assert karten[zb.id].arten == {"heizung"}
